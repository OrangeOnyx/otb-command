#!/usr/bin/env python3
"""Build PRODUCTION_TRACKER.xlsx + .csv"""
import json, csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = Path("/home/ubuntu/ON_THE_BOULEVARD")
DB = json.load(open("/home/ubuntu/tenant_database.json"))
TENANTS = DB["tenants"]
LOGO_MAP = json.load(open(ROOT / "TENANTS" / "logo_mapping.json"))

SQFT_MAP = {"101-103": 9928}
PRIORITY = {"boutique": "High", "restaurant": "High", "fitness": "Medium",
            "retail": "Medium", "salon": "Medium", "financial": "Medium",
            "services": "Low", "medical": "Low", "vacant": "Low"}

wb = Workbook()

HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

def autosize(ws):
    for col in ws.columns:
        ml = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(ml + 2, 12), 36)

# ---- Sheet 1: Master Tenant List
ws1 = wb.active
ws1.title = "Master Tenant List"
hdr1 = ["Tenant Name", "Unit Number", "Category", "Square Footage",
        "Logo Status", "Facade Family", "Illumination Type", "Priority"]
ws1.append(hdr1)
for t in TENANTS:
    has_logo = LOGO_MAP[t["tenant_name"]]["has_logo"] if t["tenant_name"] in LOGO_MAP else False
    logo_status = "On File" if has_logo else ("N/A" if t["category"] == "vacant" else "NEEDED")
    ws1.append([
        t["tenant_name"], t["unit_number"], t["category"].title(),
        SQFT_MAP.get(t["unit_number"], "TBD"),
        logo_status, t["facade_family"], t["illumination_type"] or "—",
        PRIORITY.get(t["category"], "Medium")
    ])
style_header(ws1, len(hdr1))
# Conditional formatting: Logo Status
red = PatternFill("solid", fgColor="F8CBAD")
green = PatternFill("solid", fgColor="C6EFCE")
gray = PatternFill("solid", fgColor="D9D9D9")
last = ws1.max_row
ws1.conditional_formatting.add(f"E2:E{last}", CellIsRule(operator="equal", formula=['"NEEDED"'], fill=red))
ws1.conditional_formatting.add(f"E2:E{last}", CellIsRule(operator="equal", formula=['"On File"'], fill=green))
ws1.conditional_formatting.add(f"E2:E{last}", CellIsRule(operator="equal", formula=['"N/A"'], fill=gray))
autosize(ws1)

# ---- Sheet 2: Rendering Status
ws2 = wb.create_sheet("Rendering Status")
hdr2 = ["Tenant Name", "Hero Storefront Status", "Hero Date", "Dusk Variant Status", "Dusk Date",
        "Close-Up Status", "Close-Up Date", "Streetscape Status", "Streetscape Date",
        "Revision Count", "Current Status", "Notes"]
ws2.append(hdr2)
status_options = '"Not Started,In Progress,Rendered,QC Pending,Approved,Revision,Delivered"'
for t in TENANTS:
    if t["category"] == "vacant":
        ws2.append([t["tenant_name"] + f" (Unit {t['unit_number']})", "N/A","","N/A","","N/A","","N/A","",0,"N/A",t.get("notes","")])
    else:
        ws2.append([t["tenant_name"], "Not Started","","Not Started","","Not Started","","Not Started","",0,"Not Started",""])
style_header(ws2, len(hdr2))
dv = DataValidation(type="list", formula1=status_options, allow_blank=True)
ws2.add_data_validation(dv)
for col in ["B","D","F","H","K"]:
    dv.add(f"{col}2:{col}{ws2.max_row}")
# Color code current status
for v,c in [("Approved","C6EFCE"),("Delivered","9BC2E6"),("Revision","F8CBAD"),("QC Pending","FFE699")]:
    ws2.conditional_formatting.add(f"K2:K{ws2.max_row}",
        CellIsRule(operator="equal", formula=[f'"{v}"'], fill=PatternFill("solid", fgColor=c)))
autosize(ws2)

# ---- Sheet 3: Deliverable Checklist
ws3 = wb.create_sheet("Deliverable Checklist")
hdr3 = ["Tenant Name", "Leasing Deck", "Investor Package", "Site Plan",
        "Broker Marketing", "Social Media", "Final Approval Date", "Delivered To"]
ws3.append(hdr3)
for t in TENANTS:
    if t["category"] == "vacant":
        continue
    ws3.append([t["tenant_name"], "Pending","Pending","Pending","Pending","Pending","",""])
style_header(ws3, len(hdr3))
dv2 = DataValidation(type="list", formula1='"Pending,In Progress,Complete,N/A"', allow_blank=True)
ws3.add_data_validation(dv2)
for col in ["B","C","D","E","F"]:
    dv2.add(f"{col}2:{col}{ws3.max_row}")
for col in ["B","C","D","E","F"]:
    rng = f"{col}2:{col}{ws3.max_row}"
    ws3.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Complete"'], fill=green))
    ws3.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Pending"'], fill=red))
    ws3.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"In Progress"'], fill=PatternFill("solid", fgColor="FFE699")))
autosize(ws3)

# ---- Sheet 4: QC Log
ws4 = wb.create_sheet("QC Log")
hdr4 = ["Tenant Name", "Render Type", "QC Date", "QC Result", "Issues Found", "Corrective Actions", "Approved By"]
ws4.append(hdr4)
# leave empty for operator
style_header(ws4, len(hdr4))
dv3 = DataValidation(type="list", formula1='"Pass,Fail,Conditional Pass"', allow_blank=True)
ws4.add_data_validation(dv3)
dv3.add(f"D2:D200")
dv4 = DataValidation(type="list",
    formula1='"Hero Storefront,Dusk Variant,Close-Up,Streetscape,Corner Suite,Inline"', allow_blank=True)
ws4.add_data_validation(dv4)
dv4.add(f"B2:B200")
ws4.column_dimensions['A'].width = 28
ws4.column_dimensions['E'].width = 40
ws4.column_dimensions['F'].width = 40
autosize(ws4)

# ---- Sheet 5: Production Metrics
ws5 = wb.create_sheet("Production Metrics")
total = len(TENANTS)
active = sum(1 for t in TENANTS if t["category"] != "vacant")
vacant = total - active
logos_on_file = sum(1 for v in LOGO_MAP.values() if v["has_logo"])
logos_needed = active - logos_on_file

rows = [
    ["Metric", "Value"],
    ["Total Units", total],
    ["Active Tenants", active],
    ["Vacant Units", vacant],
    ["Logos On File", logos_on_file],
    ["Logos Needed", logos_needed],
    ["Completed Renderings", "=COUNTIF('Rendering Status'!K2:K100,\"Approved\")+COUNTIF('Rendering Status'!K2:K100,\"Delivered\")"],
    ["In Progress", "=COUNTIF('Rendering Status'!K2:K100,\"In Progress\")"],
    ["QC Pending", "=COUNTIF('Rendering Status'!K2:K100,\"QC Pending\")"],
    ["Not Started", "=COUNTIF('Rendering Status'!K2:K100,\"Not Started\")"],
    ["Average Turnaround (days)", "TBD - populate after first cycle"],
    ["", ""],
    ["Tenants by Category", ""],
]
cats = {}
for t in TENANTS:
    cats[t["category"]] = cats.get(t["category"], 0) + 1
for c, n in sorted(cats.items()):
    rows.append([c.title(), n])

for r in rows:
    ws5.append(r)
style_header(ws5, 2)
ws5.column_dimensions['A'].width = 32
ws5.column_dimensions['B'].width = 22
# Bold the "Tenants by Category" header row
for row in ws5.iter_rows():
    for cell in row:
        if cell.value == "Tenants by Category":
            cell.font = Font(bold=True)

xlsx_path = ROOT / "PRODUCTION_TRACKER.xlsx"
wb.save(xlsx_path)
print("XLSX written:", xlsx_path)

# ---- CSV mirror (Master Tenant List)
with open(ROOT / "PRODUCTION_TRACKER.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(hdr1)
    for t in TENANTS:
        has_logo = LOGO_MAP[t["tenant_name"]]["has_logo"] if t["tenant_name"] in LOGO_MAP else False
        logo_status = "On File" if has_logo else ("N/A" if t["category"] == "vacant" else "NEEDED")
        w.writerow([t["tenant_name"], t["unit_number"], t["category"].title(),
                    SQFT_MAP.get(t["unit_number"], "TBD"),
                    logo_status, t["facade_family"], t["illumination_type"] or "—",
                    PRIORITY.get(t["category"], "Medium")])
print("CSV written")
